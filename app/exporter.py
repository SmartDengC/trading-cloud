from __future__ import annotations

from collections.abc import Iterable
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.sql.elements import ColumnElement

from app.config import get_settings
from app.models import DailyReview
from app.trading_service import list_trades

HEADERS = [
    "日期",
    "合约/证券代码",
    "标的",
    "市场",
    "方向",
    "策略类型",
    "周期",
    "开仓时间",
    "平仓时间",
    "持仓分钟",
    "入场理由",
    "出场理由",
    "开仓价",
    "平仓价",
    "仓位/名义金额",
    "仓位口径",
    "结算币种",
    "计划风险金额",
    "手续费税费",
    "毛盈亏",
    "净盈亏",
    "净盈亏（元）",
    "盈亏 R 倍",
    "是否盈利",
    "执行评分",
    "情绪状态",
    "错误标签",
    "错误复盘",
    "做对了什么",
    "下次改进",
    "当天交易总结",
    "截图链接",
    "状态",
    "逐笔人民币汇率",
]


def style_header(row: Iterable[Any]) -> None:
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)


async def create_export(
    db: AsyncSession, from_date: date | None = None, to_date: date | None = None
) -> bytes:
    result = await list_trades(db, from_date=from_date, to_date=to_date, page_size=500)
    review_conditions: list[ColumnElement[bool]] = [DailyReview.deleted_at.is_(None)]
    if from_date:
        review_conditions.append(DailyReview.review_date >= from_date)
    if to_date:
        review_conditions.append(DailyReview.review_date <= to_date)
    reviews = list(
        (
            await db.scalars(select(DailyReview).where(*review_conditions).order_by(DailyReview.review_date))
        ).all()
    )
    review_by_date = {item.review_date: item for item in reviews}
    workbook = Workbook()
    detail = workbook.active
    detail.title = "交易明细"
    detail.freeze_panes = "D2"
    detail.append(HEADERS)
    style_header(detail[1])
    labels = {
        "crypto": "加密",
        "a_share": "A股",
        "long": "做多",
        "short": "做空",
        "quantity": "数量",
        "notional": "金额",
        "open": "持仓中",
        "closed": "已平仓",
    }
    origin = get_settings().public_base_url
    for index, trade in enumerate(result.trades, start=2):
        review = review_by_date.get(trade.trade_date)
        detail.append(
            [
                trade.trade_date,
                trade.instrument_code,
                trade.symbol,
                labels[trade.market],
                labels[trade.side],
                trade.strategy,
                trade.timeframe,
                trade.entry_at,
                trade.exit_at,
                f'=IF(OR(H{index}="",I{index}=""),"",ROUND((I{index}-H{index})*1440,0))',
                trade.entry_reason,
                trade.exit_reason,
                float(trade.entry_price),
                float(trade.exit_price) if trade.exit_price else None,
                float(trade.position_size),
                labels[trade.position_basis],
                trade.settlement_currency,
                float(trade.planned_risk_amount) if trade.planned_risk_amount else None,
                float(trade.fees or 0),
                float(trade.gross_pnl) if trade.gross_pnl else None,
                float(trade.net_pnl) if trade.net_pnl else None,
                float(trade.pnl_cny) if trade.pnl_cny else None,
                float(trade.r_multiple) if trade.r_multiple else None,
                "" if trade.is_winning is None else "是" if trade.is_winning else "否",
                trade.execution_grade,
                trade.emotion,
                "、".join(trade.error_tags),
                trade.error_notes,
                trade.did_well,
                trade.next_improvement,
                review.daily_summary if review else None,
                "\n".join(f"{origin}{item.file_url}" for item in trade.attachments),
                labels[trade.status],
                float(trade.fx_to_cny),
            ]
        )
    detail.auto_filter.ref = f"A1:AH{max(detail.max_row, 1)}"
    widths = [
        12,
        16,
        24,
        10,
        10,
        16,
        10,
        19,
        19,
        12,
        36,
        30,
        14,
        14,
        18,
        12,
        12,
        16,
        14,
        14,
        14,
        16,
        12,
        12,
        12,
        12,
        24,
        36,
        32,
        32,
        36,
        48,
        12,
        16,
    ]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[detail.cell(1, index).column_letter].width = width
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    daily = workbook.create_sheet("每日复盘")
    daily.append(
        [
            "日期",
            "市场环境 / 盘前计划",
            "每日总结",
            "最大失误与原因",
            "明日只改一件事",
            "只做计划内交易",
            "严格执行止损",
            "无临盘加仓冲动",
            "无报复性交易",
            "按计划离场",
            "优先修正项",
            "备注",
        ]
    )
    style_header(daily[1])

    def yes_no(value: bool | None) -> str:
        return "" if value is None else "是" if value else "否"

    for review in reviews:
        daily.append(
            [
                review.review_date,
                review.market_plan,
                review.daily_summary,
                review.biggest_mistake,
                review.tomorrow_one_thing,
                yes_no(review.planned_only),
                yes_no(review.followed_stops),
                yes_no(review.avoided_impulse_adds),
                yes_no(review.avoided_revenge_trading),
                yes_no(review.exited_as_planned),
                review.priority_fix,
                review.notes,
            ]
        )

    summary = workbook.create_sheet("统计摘要")
    summary.append(["私有交易复盘 · 导出摘要"])
    summary.append([])
    summary.append(["指标", "数值"])
    summary.append(["交易总数", len(result.trades)])
    summary.append(["已平仓交易", sum(item.status == "closed" for item in result.trades)])
    summary.append(["净盈亏（元）", sum(float(item.pnl_cny or 0) for item in result.trades)])
    summary.append(["盈利笔数", sum(bool(item.is_winning) for item in result.trades)])
    closed = sum(item.status == "closed" for item in result.trades)
    summary.append(
        ["胜率", sum(bool(item.is_winning) for item in result.trades) / closed if closed else None]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
